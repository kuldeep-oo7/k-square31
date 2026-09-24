"""
THE FUTURES WATCH LIST IS YOUR WORKBOOK: ALL 288 ROWS.

Asked for, verbatim (2026-09-23): "in the futures I need you from the excel file
we are using ... you need to give me only from that", then, pointing at
NSE_FO_Symbols_Auto_Update.xlsm: "this 288 you need to watch, understand".

The workbook's "All Symbols" sheet has 288 rows: 3 indices and 285 stocks. This
reads it exactly -- symbol, type and the Yahoo ticker from "Live Prices" -- and
writes data/watchlist.json, which the futures engine uses in place of NSE's own
F&O list.

What the desk then does with each name is decided by the exchange, not by this
file, and the page says which is which:
  - a name NSE still lists in F&O has a futures contract, priced live;
  - a name NSE removed from F&O has no futures contract at all, and is watched
    on its share price;
  - a symbol NSE does not list under that name (a merger, a rename, a typo) is
    named on the page as unpriceable, never dropped silently.

Run it whenever the workbook changes:

    python scripts/import-watchlist-excel.py "C:/Users/kulde/Downloads/NSE_FO_Symbols_Auto_Update.xlsm"
"""

import datetime
import json
import os
import sys

import openpyxl

SOURCE = sys.argv[1] if len(sys.argv) > 1 else r"C:/Users/kulde/Downloads/NSE_FO_Symbols_Auto_Update.xlsm"
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "watchlist.json")

wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True, keep_vba=False)

yahoo = {}
for row in wb["Live Prices"].iter_rows(min_row=2, values_only=True):
    cells = list(row) + [None] * 4
    symbol, ticker = cells[1], cells[3]
    if symbol and ticker:
        yahoo[str(symbol).strip().upper()] = str(ticker).strip()

indices, stocks = [], []
for row in wb["All Symbols"].iter_rows(min_row=2, values_only=True):
    cells = list(row) + [None] * 3
    symbol, kind = cells[1], cells[2]
    if not symbol:
        continue
    entry = {"symbol": str(symbol).strip().upper(), "yahoo": yahoo.get(str(symbol).strip().upper())}
    (indices if str(kind).strip().lower() == "index" else stocks).append(entry)

modified = datetime.datetime.fromtimestamp(os.path.getmtime(SOURCE)).strftime("%Y-%m-%d")
payload = {
    "source": os.path.basename(SOURCE) + ", sheet All Symbols (Yahoo tickers from Live Prices)",
    "capturedAt": datetime.date.today().isoformat(),
    "workbookModified": modified,
    "note": ("The futures watch list (user instruction, 2026-09-23: 'this 288 you need to watch'). "
             "Options are not limited by it: they come from every F&O name the exchanges list."),
    "indices": indices,
    "stocks": stocks,
}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(payload, f, indent=2, ensure_ascii=False)
    f.write("\n")
print(f"{len(indices)} indices + {len(stocks)} stocks = {len(indices) + len(stocks)} from "
      f"{os.path.basename(SOURCE)} (modified {modified}) -> data/watchlist.json")
