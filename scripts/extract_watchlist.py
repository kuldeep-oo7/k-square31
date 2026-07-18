"""Convert the source workbook into the browser's read-only fallback snapshot."""

from pathlib import Path
import csv
import json
import openpyxl

SOURCE = Path(r"C:\Users\kulde\Downloads\journal (1).xlsx")
TARGET = Path(__file__).resolve().parents[1] / "watchlist-data.js"
DATA_DIR = Path(__file__).resolve().parents[1] / "data"

formula_book = openpyxl.load_workbook(SOURCE, data_only=False, read_only=True)
value_book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
formula_rows = list(formula_book["SCREENER"].iter_rows(min_row=2, values_only=True))
value_rows = list(value_book["SCREENER"].iter_rows(min_row=2, values_only=True))

amfi_book = openpyxl.load_workbook(DATA_DIR / "amfi-market-cap-2026H1.xlsx", data_only=True, read_only=True)
amfi_rows = amfi_book["FINAL"].iter_rows(min_row=3, values_only=True)
market_caps = {
    str(row[5]).strip(): {
        "marketCapRank": row[0],
        "marketCapCrore": row[9],
        "marketCapCategory": row[10],
    }
    for row in amfi_rows
    if len(row) > 10 and row[5] not in (None, "-")
}

with (DATA_DIR / "nifty500-constituents.csv").open(encoding="utf-8-sig", newline="") as handle:
    sectors = {row["Symbol"].strip(): row["Industry"].strip() for row in csv.DictReader(handle)}
sectors.update({"STAR": "Healthcare", "APLLTD": "Healthcare"})

items = []
for formula_row, value_row in zip(formula_rows, value_rows):
    symbol = formula_row[2] if len(formula_row) > 2 else None
    if not symbol:
        continue
    symbol = str(symbol).strip()
    cap = market_caps.get(symbol, {})
    items.append(
        {
            "symbol": symbol,
            "exchange": "NSE",
            "sector": sectors.get(symbol, "Other"),
            "marketCapCategory": cap.get("marketCapCategory", "Unclassified"),
            "marketCapRank": cap.get("marketCapRank"),
            "marketCapCrore": cap.get("marketCapCrore"),
            "snapshotPrice": value_row[3],
            "snapshotDayHigh": value_row[4],
            "snapshotDayLow": value_row[5],
            "snapshotChangePct": value_row[6],
            "snapshot52High": value_row[8],
            "snapshot52Low": value_row[9],
        }
    )

payload = "window.WATCHLIST = " + json.dumps(items, indent=2) + ";\n"
TARGET.write_text(payload, encoding="utf-8")
print(f"Wrote {len(items)} symbols to {TARGET}")
